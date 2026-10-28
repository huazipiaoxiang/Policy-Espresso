from pathlib import Path
import shutil
import subprocess
import tempfile


class UploadParseError(ValueError):
    pass


def extract_upload(path: Path, suffix: str) -> str:
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
    if suffix == ".docx":
        from docx import Document as DocxDocument

        document = DocxDocument(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables for row in table.rows
        )
        return "\n".join(parts).strip()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(
                "\t".join("" if value is None else str(value) for value in row)
                for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)
            ).strip()
        finally:
            workbook.close()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(path)
        return "\n".join(
            "\t".join(str(value) for value in sheet.row_values(row_index))
            for sheet in workbook.sheets() for row_index in range(sheet.nrows)
        ).strip()
    if suffix == ".doc":
        return _extract_legacy_doc(path)
    raise UploadParseError("仅支持 PDF、DOC、DOCX、XLS、XLSX 文件")


def _extract_legacy_doc(path: Path) -> str:
    antiword = shutil.which("antiword")
    if antiword:
        result = subprocess.run(
            [antiword, str(path)], capture_output=True, text=True, timeout=60, check=False,
        )
        if result.returncode != 0:
            raise UploadParseError(f"DOC 解析失败：{result.stderr.strip() or '无法读取文件'}")
        return result.stdout.strip()

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory(prefix="horizon-doc-") as output_dir:
            result = subprocess.run(
                [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", output_dir, str(path)],
                capture_output=True, text=True, timeout=90, check=False,
            )
            output = Path(output_dir) / f"{path.stem}.txt"
            if result.returncode == 0 and output.exists():
                return output.read_text(encoding="utf-8", errors="replace").strip()
            raise UploadParseError(f"DOC 解析失败：{result.stderr.strip() or '无法读取文件'}")

    raise UploadParseError("服务器未安装 DOC 解析器（antiword 或 LibreOffice）；原文件已保存")
